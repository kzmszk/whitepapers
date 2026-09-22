"""Extract published wage growth rates without crossing index discontinuities."""
import csv
import hashlib
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "raw/context/wages-2025-final.xlsx"


def main():
    sheet = openpyxl.load_workbook(SOURCE, data_only=True, read_only=True)["賃金指数"]
    assert sheet["D6"].value == "前年比"
    assert sheet["E6"].value == "実質前年比"
    assert sheet["A8"].value.replace("\u3000", "") == "現金給与総額"
    observations = []
    for row, year in zip(range(9, 13), range(2022, 2026)):
        assert sheet.cell(row, 1).value.strip() == f"{year}年"
        for column, name in [(4, "nominal"), (5, "real")]:
            value = sheet.cell(row, column).value
            assert isinstance(value, (int, float))
            observations.append({"year": year, "series_id": name, "value": value,
                                 "unit": "percent_yoy", "source_cell": sheet.cell(row, column).coordinate,
                                 "source_id": "wages-mhlw-timeseries"})
    expected = {"nominal": [2.0, 1.2, 2.8, 2.3], "real": [-1.0, -2.5, -0.3, -1.3]}
    series_values = {name: [row["value"] for row in observations if row["series_id"] == name] for name in expected}
    assert series_values == expected
    chart = {
        "id": "nominal-real-wages", "title": "給料の額が増えても、買える量は増えていない",
        "question": "物価の上昇を差し引いても賃金は増えた？", "type": "line",
        "unit": "前年比（%）", "period_type": "calendar_year", "period_label": "暦年平均",
        "periods": [str(year) for year in range(2022, 2026)],
        "series": [{"id": name, "label": label, "values": series_values[name],
                    "source_ids": ["wages-mhlw-timeseries"]}
                   for name, label in [("nominal", "名目賃金"), ("real", "実質賃金")]],
        "notes": ["調査産業計・事業所規模5人以上の常用労働者。中小企業だけの値ではない。",
                  "現金給与総額は賞与等を含む。実質は消費者物価指数『持家の帰属家賃を除く総合』で調整。",
                  "2024年のベンチマーク更新による断層を避けるため、公表された前年比を使用。指数の単純な割り算ではない。",
                  "2025年12月確報に掲載された年平均値を固定して使用。各人の手取りや同一人物の昇給率ではない。"],
        "derivation": "毎月勤労統計・時系列表第1表『賃金指数』A9:A12を年、D9:D12を名目前年比、E9:E12を実質前年比として抽出。再計算・補間なし。",
        "takeaway": "2025年の名目賃金は2.3%増えたが、物価を考慮した実質賃金は1.3%減った。2022〜2025年の実質前年比はいずれもマイナス。"
    }
    charts = [chart]
    (ROOT / "processed/wages-charts.json").write_text(json.dumps(charts, ensure_ascii=False, indent=2) + "\n")
    with (ROOT / "processed/wages-observations.csv").open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=list(observations[0]))
        writer.writeheader(); writer.writerows(observations)
    checks = {"source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
              "checks": "8 published annual growth rates matched to source cells and independently inspected PDF table",
              "observations": observations, "uses_published_growth_rates": True}
    (ROOT / "processed/wages-checks.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2) + "\n")
    article = ROOT / "content/wages.json"
    if article.exists():
        content = json.loads(article.read_text()); content["charts"] = charts
        article.write_text(json.dumps(content, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"charts": len(charts), "observations": len(observations)}))


if __name__ == "__main__":
    main()
