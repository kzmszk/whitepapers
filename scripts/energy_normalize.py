#!/usr/bin/env python3
"""Normalize official fiscal-year energy balances; never infer IEA self-sufficiency from HCV totals."""
import csv, json, math, pathlib, re
import openpyxl
ROOT=pathlib.Path(__file__).resolve().parents[1]
OUT=ROOT/'processed'; OUT.mkdir(exist_ok=True)
SERIES=[
 ('primary_supply_total','一次エネルギー国内供給','一次エネルギー供給','#01','#100000','$1400'),
 ('domestic_production_total','一次エネルギー国内産出（HCV）','国内産出','#02','#110000','$1400'),
 ('primary_supply_coal','一次供給：石炭','一次エネルギー供給','#01','#100000','$0100'),
 ('primary_supply_coal_products','一次供給：石炭製品','一次エネルギー供給','#01','#100000','$0200'),
 ('primary_supply_crude_oil','一次供給：原油','一次エネルギー供給','#01','#100000','$0300'),
 ('primary_supply_oil_products','一次供給：石油製品','一次エネルギー供給','#01','#100000','$0400'),
 ('primary_supply_natural_gas','一次供給：天然ガス','一次エネルギー供給','#01','#100000','$0500'),
 ('primary_supply_city_gas','一次供給：都市ガス','一次エネルギー供給','#01','#100000','$0600'),
 ('primary_supply_renewables_excl_hydro','一次供給：再生可能エネルギー（水力除く）','一次エネルギー供給','#01','#100000','$0700'),
 ('primary_supply_hydro','一次供給：水力（揚水除く）','一次エネルギー供給','#01','#100000','$0800'),
 ('primary_supply_pumped','一次供給：揚水','一次エネルギー供給','#01','#100000','$0900'),
 ('primary_supply_recovered','一次供給：未活用エネルギー','一次エネルギー供給','#01','#100000','$1000'),
 ('primary_supply_nuclear','一次供給：原子力','一次エネルギー供給','#01','#100000','$1100'),
 ('final_consumption_total','最終エネルギー消費：合計','最終エネルギー消費','#19','#500000','$1400'),
 ('final_consumption_industry_business','最終消費：企業・事業所他','企業・事業所他','#20','#600000','$1400'),
 ('final_consumption_manufacturing','最終消費：製造業','製造業','#22','#620000','$1400'),
 ('final_consumption_commercial','最終消費：業務他','業務他','#32','#650000','$1400'),
 ('final_consumption_households','最終消費：家庭','家庭','#33','#700000','$1400'),
 ('final_consumption_transport','最終消費：運輸','運輸','#34','#800000','$1400'),
 ('final_consumption_electricity','最終消費：電力','最終エネルギー消費','#19','#500000','$1200'),
]
records=[]
for year in range(1990,2025):
    path=ROOT/'raw'/'energy'/f'energy-balance-{year}-simple.xlsx'
    book=openpyxl.load_workbook(path,data_only=True,read_only=True)
    sheet=next(s for s in book if 'ｴﾈﾙｷﾞｰ単位表' in s.title)
    cells=list(sheet.iter_rows(max_row=min(sheet.max_row,900),max_col=min(sheet.max_column,180)))
    assert str(cells[0][0].value).strip()==f'{year}FY'
    cols={c.value:i for i,c in enumerate(cells[0]) if str(c.value).startswith('$')}
    rowmap={}
    for row in cells:
        if isinstance(row[0].value,str) and row[0].value.startswith('#'):
            rowmap.setdefault(row[0].value,row)
    detail='詳細' in sheet.title
    values={}
    for series,label,rowlabel,simple,full,column in SERIES:
        row=rowmap[full if detail else simple]
        assert str(row[3].value).strip()==rowlabel,(year,series,row[3].value)
        col=cols[column];assert cells[13][col].value=='TJ'
        cell=row[col];value=cell.value
        assert isinstance(value,(int,float)) and math.isfinite(value)
        values[series]=value/1000
        records.append(dict(year=year,period=f'FY{year}',period_type='fiscal_year_april_march',geography='JP',series=series,label=label,value=value/1000,unit='PJ',energy_basis='HCV',source_id=f'energy-balance-{year}',source_sheet=sheet.title,source_cell=cell.coordinate,source_unit='TJ',transformation='value / 1000',revision_date='2026-04-14'))
    assert math.isclose(values['final_consumption_total'],sum(values[x] for x in ['final_consumption_industry_business','final_consumption_households','final_consumption_transport']),rel_tol=1e-8)
    components=[v for k,v in values.items() if k.startswith('primary_supply_') and k!='primary_supply_total']
    assert math.isclose(values['primary_supply_total'],sum(components),rel_tol=1e-8)
    book.close()
with (OUT/'energy-annual.csv').open('w',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=list(records[0]));writer.writeheader();writer.writerows(records)
(OUT/'energy-annual.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n')
text=(ROOT/'raw'/'energy'/'energy-supply-demand-2024-summary.txt').read_text()
line=next(s for s in text.splitlines() if s.startswith('エネルギー自給率(IEAベース)'))
vals=[float(n) for n in re.findall(r'([\d.]+)%',line)]
assert len(vals)==15 and vals[-1]==16.3
with (OUT/'energy-self-sufficiency-published.csv').open('w',newline='') as f:
    writer=csv.writer(f);writer.writerow(['year','period_type','geography','value','unit','basis','source_id','source_page','transformation','precision'])
    writer.writerows([y,'fiscal_year_april_march','JP',v,'percent','IEA_basis','energy-supply-demand-2024-summary',7,'PDF text table extraction; not recalculation','published 1 decimal'] for y,v in zip(range(2010,2025),vals))
summary={
 'records':len(records),'years':35,'series':len(SERIES),'range':[1990,2024],
 'checks':['Workbook fiscal year and row labels verified for every series/year','TJ units verified and converted to PJ by /1000','Final consumption equals enterprise/business + households + transport for every year','Primary supply component sum equals primary supply total for every year','PDF self-sufficiency 15 values FY2010–FY2024, last value visually verified 16.3%'],
 'source_anomalies':['e-Stat links labeled 1994年度簡易表 and 1995年度簡易表 return 詳細表 workbooks. They are preserved unchanged and extracted using row/column codes, not fixed cell addresses.'],
 'caveats':['HCV balance-table domestic production / primary supply is not the published IEA-basis self-sufficiency ratio. The latter uses a different statistical basis; use the separate published series, or obtain the official lower-calorific IEA reference workbook to reproduce.','Final consumption total includes non-energy use; do not relabel as only fuel burned.','Primary supply mix is not electricity generation mix.','Fiscal year is April through following March; calendar-year imports/CPI cannot be joined as identical periods.','Standard calorific values were updated from FY2023; consider this when comparing physical quantities across years.']}
(OUT/'energy-validation.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(summary,ensure_ascii=False))
